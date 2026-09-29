import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from core import annotate_pdf, export_table, extract_requirements, suggest_documents


class CoreTests(unittest.TestCase):
    def test_pdf_import_keeps_clause_numbers_and_physical_pages(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "tor.pdf"
            pdf = canvas.Canvas(str(source))
            pdf.drawString(50, 750, "4.15.2 IPv6 support")
            pdf.drawString(50, 730, "for the network")
            pdf.showPage()
            pdf.drawString(50, 750, "4.16 NOC report")
            pdf.save()

            rows, unreadable = extract_requirements(source)

            self.assertEqual([r["number"] for r in rows], ["4.15.2", "4.16"])
            self.assertEqual([r["source_page"] for r in rows], [1, 2])
            self.assertIn("for the network", rows[0]["text"])
            self.assertEqual(unreadable, [])

    def test_pdf_import_preserves_continuation_before_next_page_clause(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "tor.pdf"
            pdf = canvas.Canvas(str(source))
            pdf.drawString(50, 750, "4.15.2 Must provide IPv6")
            pdf.showPage()
            pdf.drawString(50, 750, "including network monitoring")
            pdf.drawString(50, 720, "4.16 Must provide NOC")
            pdf.save()

            rows, _ = extract_requirements(source)

            self.assertIn("including network monitoring", rows[0]["text"])
            self.assertEqual(rows[1]["number"], "4.16")

    def test_scanned_pdf_page_is_reported_for_manual_review(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "scan.pdf"
            pdf = canvas.Canvas(str(source))
            pdf.rect(50, 50, 100, 100)
            pdf.save()

            rows, unreadable = extract_requirements(source)

            self.assertEqual(rows, [])
            self.assertEqual(unreadable, [1])

    def test_docx_import_reads_numbered_table_rows(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "tor.docx"
            doc = Document()
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "ข้อ"
            table.cell(0, 1).text = "รายละเอียด"
            table.cell(1, 0).text = "5.3"
            table.cell(1, 1).text = "รองรับ IPv6 Addressing"
            doc.save(source)

            rows, unreadable = extract_requirements(source)

            self.assertEqual(unreadable, [])
            self.assertEqual(rows[0]["number"], "5.3")
            self.assertIn("IPv6", rows[0]["text"])

    def test_ocr_parser_recovers_number_after_short_misread_prefix(self):
        from core import parse_text_page

        rows = parse_text_page("ข้� 5.3 รองรับ IPv6 Addressing", 1)
        self.assertEqual(rows[0]["number"], "5.3")

    def test_ocr_parser_preserves_leading_continuation_text(self):
        from core import parse_text_page

        rows = parse_text_page("including network monitoring\n4.16 Must provide NOC", 2)
        self.assertIn("including network monitoring", rows[0]["text"])
        self.assertEqual(rows[1]["number"], "4.16")

    def test_suggestions_rank_exact_product_and_document_terms_without_claiming_compliance(self):
        docs = [
            {"id": 1, "filename": "network.pdf", "product_name": "Router A", "text": "MPLS IPv6 routing"},
            {"id": 2, "filename": "storage.pdf", "product_name": "NAS B", "text": "disk storage"},
        ]
        ranked = suggest_documents("ต้องรองรับ IPv6 สำหรับ MPLS", docs)
        self.assertEqual(ranked[0]["id"], 1)
        self.assertIn("IPv6", ranked[0]["matched_terms"])
        self.assertNotIn("compliant", ranked[0])

    def test_export_has_four_columns_and_uses_printed_page_reference(self):
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "comply.xlsx"
            export_table(
                target,
                [{"number": "5.3", "text": "รองรับ IPv6", "proposal": "เสนอ Router A", "comparison": "ตรงตามข้อกำหนด", "evidence": [{"filename": "technical.pdf", "pdf_page": 5, "printed_page": "4"}]}],
            )
            sheet = load_workbook(target).active
            self.assertEqual(sheet.max_column, 4)
            self.assertEqual(sheet["C2"].value, "ตรงตามข้อกำหนด")
            self.assertIn("หน้า 4", sheet["D2"].value)
            self.assertIn("PDF 5", sheet["D2"].value)

    def test_excel_export_keeps_formula_like_text_literal_and_rows_auto_height(self):
        with tempfile.TemporaryDirectory() as temp:
            target = Path(temp) / "comply.xlsx"
            proposal = '=HYPERLINK("https://example.com","bad")'
            export_table(target, [{"number": "1.1", "text": "very long requirement " * 30, "proposal": proposal, "evidence": []}])
            sheet = load_workbook(target).active
            self.assertEqual(sheet["B2"].value, proposal)
            self.assertEqual(sheet["B2"].data_type, "s")
            self.assertIsNone(sheet.row_dimensions[2].height)

    def test_marked_pdf_contains_clause_label_and_keeps_original_page(self):
        with tempfile.TemporaryDirectory() as temp:
            source = Path(temp) / "evidence.pdf"
            target = Path(temp) / "marked.pdf"
            pdf = canvas.Canvas(str(source))
            pdf.drawString(50, 700, "IPv6 Addressing")
            pdf.save()
            annotate_pdf(source, target, [{"pdf_page": 1, "number": "5.3", "box": [0.07, 0.10, 0.30, 0.04]}])
            reader = PdfReader(target)
            self.assertEqual(len(reader.pages), 1)
            self.assertIn("IPv6 Addressing", reader.pages[0].extract_text())
            self.assertIn("5.3", reader.pages[0].extract_text())


if __name__ == "__main__":
    unittest.main()
