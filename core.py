"""Document operations for the local TOR comply workflow."""

from __future__ import annotations

import io
import re
from pathlib import Path

import pdfplumber
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from pypdf import PdfReader, PdfWriter
from reportlab.lib import colors
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


CLAUSE = re.compile(r"^\s*(?:ข้อ\s*)?(\d+(?:\.\d+)*)(?:[.)])?\s+(.{3,})$")
TERMS = ("MPLS", "IPv6", "IPv4", "NOC", "SLA", "VPN", "Firewall", "Switch", "Router", "Server", "Storage", "Network Monitor", "วงจร", "เครือข่าย", "รายงาน", "บริการ")
HEADERS = (
    "รายละเอียดการดำเนินงาน",
    "รายละเอียดการดำเนินงานที่ผู้เสนอราคาเสนอ",
    "เปรียบเทียบรายละเอียดการดำเนินงานที่ผู้เสนอราคาเสนอ",
    "เอกสารอ้างอิง ไฟล์ใด หน้าใด",
)


def _clause_match(raw: str):
    line = " ".join(raw.split())
    match = CLAUSE.match(line)
    return match if match and ("." in match.group(1) or raw.lstrip().startswith("ข้อ") or len(line) > 25) else None


def _leading_lines(lines: list[str]) -> list[str]:
    for index, line in enumerate(lines):
        if _clause_match(line):
            return lines[:index]
    return lines


def _numbered_lines(lines: list[str], page: int | None) -> list[dict]:
    rows: list[dict] = []
    for raw in lines:
        line = " ".join(raw.split())
        if not line:
            continue
        match = _clause_match(raw)
        if match:
            rows.append({"number": match.group(1), "text": match.group(2), "source_page": page})
        elif rows:
            rows[-1]["text"] += " " + line
    return rows


def parse_text_page(text: str, page: int) -> list[dict]:
    """Turn one OCR page into editable candidate clauses."""
    cleaned = [re.sub(r"^\D{1,12}(?=\d+(?:\.\d+)+\s)", "", line) for line in text.splitlines()]
    found = _numbered_lines(cleaned, page)
    if found:
        leading = " ".join(" ".join(line.split()) for line in _leading_lines(cleaned)).strip()
        return ([{"number": f"หน้า {page} (ข้อความต่อเนื่อง)", "text": leading, "source_page": page}] if leading else []) + found
    return [{"number": f"หน้า {page}", "text": " ".join(text.split()), "source_page": page}] if text.strip() else []


def extract_requirements(path: str | Path) -> tuple[list[dict], list[int]]:
    """Extract candidate TOR clauses; report image-only PDF pages explicitly."""
    path = Path(path)
    if path.suffix.lower() == ".pdf":
        rows: list[dict] = []
        unreadable: list[int] = []
        with pdfplumber.open(path) as pdf:
            for page_number, page in enumerate(pdf.pages, 1):
                text = page.extract_text() or ""
                if not text.strip():
                    unreadable.append(page_number)
                    continue
                lines = text.splitlines()
                found = _numbered_lines(lines, page_number)
                if found:
                    leading = " ".join(" ".join(line.split()) for line in _leading_lines(lines)).strip()
                    if leading:
                        if rows:
                            rows[-1]["text"] += f" [ต่อหน้า PDF {page_number}] " + leading
                        else:
                            rows.append({"number": f"หน้า {page_number} (ข้อความก่อนเลขข้อ)", "text": leading, "source_page": page_number})
                    rows.extend(found)
                else:
                    if rows:
                        rows[-1]["text"] += f" [ต่อหน้า PDF {page_number}] " + " ".join(text.split())
                    else:
                        rows.append({"number": f"หน้า {page_number}", "text": " ".join(text.split()), "source_page": page_number})
        return rows, unreadable
    if path.suffix.lower() == ".docx":
        doc = Document(path)
        rows = _numbered_lines([p.text for p in doc.paragraphs], None)
        for table in doc.tables:
            for tr in table.rows:
                cells = [" ".join(c.text.split()) for c in tr.cells]
                if len(cells) < 2:
                    continue
                number = cells[0].strip().rstrip(".)")
                if re.fullmatch(r"\d+(?:\.\d+)*", number) and cells[1]:
                    rows.append({"number": number, "text": cells[1], "source_page": None})
        return rows, []
    raise ValueError("รองรับ TOR แบบ PDF หรือ DOCX เท่านั้น")


def suggest_documents(requirement: str, documents: list[dict]) -> list[dict]:
    """Rank uploaded documents by literal term overlap; never infer compliance."""
    query = requirement.casefold()
    terms = {term for term in TERMS if term.casefold() in query}
    terms.update(re.findall(r"[A-Za-z][A-Za-z0-9._+-]{2,}", requirement))
    ranked = []
    for doc in documents:
        haystack = " ".join(str(doc.get(key, "")) for key in ("filename", "product_name", "text")).casefold()
        matched = sorted((term for term in terms if term.casefold() in haystack), key=str.casefold)
        if matched:
            ranked.append({**doc, "matched_terms": matched, "match_count": len(matched)})
    return sorted(ranked, key=lambda item: (-item["match_count"], item["filename"].casefold()))


def export_table(path: str | Path, rows: list[dict]) -> None:
    """Write the four-column comply table without inventing a pass judgment."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ตาราง Comply TOR"
    sheet.append(HEADERS)
    for row in rows:
        references = []
        for item in row.get("evidence", []):
            printed = str(item.get("printed_page") or "").strip()
            page = int(item["pdf_page"])
            page_text = f"หน้า {printed} (PDF {page})" if printed else f"หน้า PDF {page}"
            references.append(f"{item['filename']} {page_text}")
        requirement = f"ข้อ {row['number']} {row['text']}".strip()
        sheet.append((requirement, row.get("proposal") or "", row.get("comparison") or "รอตรวจสอบ", "\n".join(references)))
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:D{max(2, sheet.max_row)}"
    for col, width in zip("ABCD", (48, 48, 33, 43)):
        sheet.column_dimensions[col].width = width
    red = "FF0038"
    thin = Side(style="thin", color="D9D9DD")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=red)
        cell.font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Tahoma", size=10, color="262629")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
    workbook.save(path)


def _annotation_font() -> str:
    font_path = Path("C:/Windows/Fonts/tahoma.ttf")
    if font_path.exists():
        if "TORThai" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("TORThai", str(font_path)))
        return "TORThai"
    return "Helvetica"


def annotate_pdf(source: str | Path, target: str | Path, marks: list[dict]) -> None:
    """Flatten yellow highlights and visible red clause labels into a PDF copy."""
    reader = PdfReader(source)
    writer = PdfWriter()
    font = _annotation_font()
    for index, source_page in enumerate(reader.pages, 1):
        page = writer.add_page(source_page)
        page_marks = [mark for mark in marks if int(mark["pdf_page"]) == index]
        if page_marks:
            width = float(page.mediabox.width)
            height = float(page.mediabox.height)
            overlay = io.BytesIO()
            layer = canvas.Canvas(overlay, pagesize=(width, height))
            for mark in page_marks:
                x, top, box_width, box_height = [float(v) for v in mark["box"]]
                if not (0 <= x < 1 and 0 <= top < 1 and 0 < box_width <= 1 - x and 0 < box_height <= 1 - top):
                    raise ValueError("ตำแหน่งไฮไลต์อยู่นอกหน้าเอกสาร")
                left = x * width
                bottom = (1 - top - box_height) * height
                layer.setFillColor(colors.HexColor("#FFF176"))
                layer.setFillAlpha(0.48)
                layer.rect(left, bottom, box_width * width, box_height * height, fill=1, stroke=0)
                layer.setFillAlpha(1)
                layer.setFillColor(colors.HexColor("#D70030"))
                layer.setFont(font, 9)
                label = f"ข้อที่ {mark['number']}" if font == "TORThai" else f"Clause {mark['number']}"
                label_y = min(height - 11, bottom + box_height * height + 3)
                layer.drawString(left, label_y, label)
            layer.save()
            overlay.seek(0)
            page.merge_page(PdfReader(overlay).pages[0])
    with Path(target).open("wb") as stream:
        writer.write(stream)
